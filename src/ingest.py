"""
Parse every source document into a list of (page_text, page_number, doc_id)
records, with an OCR fallback for scanned PDFs.

Run directly to see per-document parsing stats:
    python -m src.ingest
"""
from __future__ import annotations

import json
import sys
from dataclasses import dataclass, asdict
from pathlib import Path

import pymupdf
from docx import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from src.config import DOCUMENTS, DATA_PROCESSED  # noqa: E402


@dataclass
class Page:
    doc_id: str
    title: str
    doc_type: str
    effective_date: str
    supersedes: list
    page_number: int
    text: str
    ocr: bool = False
    faculty: str | None = None
    approx_pages: bool = False


FIXED_PAGE_CHARS = 3500  # matches the observed density of a real handbook page


def _ocr_page(pdf_path: Path, page_number: int) -> str:
    """OCR a single page. Imported lazily -- only the ragging circular needs it."""
    from pdf2image import convert_from_path
    import pytesseract

    images = convert_from_path(
        str(pdf_path), first_page=page_number, last_page=page_number, dpi=300
    )
    if not images:
        return ""
    return pytesseract.image_to_string(images[0])


def _fixed_size_pseudo_pages(raw: str, size: int = FIXED_PAGE_CHARS) -> list[str]:
    """
    Fallback for a .txt export with no detectable repeating page-footer
    (see fmas_handbook_2025 in config.py). Breaks on paragraph boundaries
    near the target size rather than mid-sentence, so citation units stay
    readable -- but the resulting page numbers are frankly a citation
    convenience, not the original PDF's real pagination. Documented, not
    hidden: see error_analysis.md finding #5.
    """
    paras = [p for p in raw.split("\n\n") if p.strip()]
    pages, buf, buf_len = [], [], 0
    for para in paras:
        if buf_len + len(para) > size and buf:
            pages.append("\n\n".join(buf))
            buf, buf_len = [], 0
        buf.append(para)
        buf_len += len(para)
    if buf:
        pages.append("\n\n".join(buf))
    return pages


def parse_txt(doc: dict) -> list[Page]:
    """
    Two of the source documents arrived as pre-extracted .txt files
    rather than PDFs. We still need page-like units for citations:

    - If the export has a reliably repeating footer string (set as
      `page_footer` in config.py), split on it -- this tracks the
      original PDF's real page breaks closely (see
      student_handbook_2022_23).
    - Otherwise (fmas_handbook_2025 has no such marker), fall back to
      fixed-size pseudo-pages. This is a parsing decision that has to be
      justified rather than assumed -- see error_analysis.md.
    """
    raw = Path(doc["path"]).read_text(encoding="utf-8", errors="ignore")
    footer = doc.get("page_footer")
    if footer and raw.count(footer) >= 5:
        raw_pages = raw.split(footer)
        approx = False
    else:
        raw_pages = _fixed_size_pseudo_pages(raw)
        approx = True

    pages = []
    for i, chunk in enumerate(raw_pages, start=1):
        text = chunk.strip()
        if text:
            pages.append(
                Page(
                    doc_id=doc["doc_id"],
                    title=doc["title"],
                    doc_type=doc["doc_type"],
                    effective_date=doc["effective_date"],
                    supersedes=doc["supersedes"],
                    page_number=i,
                    text=text,
                    faculty=doc.get("faculty"),
                    approx_pages=approx or doc.get("approx_pages", False),
                )
            )
    return pages


def _iter_block_items(docx_doc):
    """
    Walk a .docx body in document order, yielding Paragraph and Table
    objects as they actually appear. python-docx's own `.paragraphs`
    convenience property SKIPS any paragraph that lives inside a table
    cell -- for the Management Studies handbook that meant the whole
    "Evaluation Procedure" / "Examination Offences & Punishments"
    section could have silently vanished if it had been table-formatted
    (it turned out that section wasn't in this export at all -- see
    error_analysis.md finding #11 -- but the earlier Agriculture/FMAS
    documents made clear that trusting `.paragraphs` alone is not safe
    in general).
    """
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P

    for child in docx_doc.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, docx_doc)
        elif isinstance(child, CT_Tbl):
            yield Table(child, docx_doc)


def parse_docx(doc: dict) -> list[Page]:
    """
    Two handbooks arrived as .docx exports (smaller than the PDF
    originals). Extract paragraphs AND table cells, in document order,
    then fall back to the same fixed-size pseudo-paging used for the
    FMAS .txt export -- a .docx has no native page boundaries either.
    """
    docx_doc = DocxDocument(str(doc["path"]))
    parts = []
    for block in _iter_block_items(docx_doc):
        if isinstance(block, Paragraph):
            if block.text.strip():
                parts.append(block.text.strip())
        elif isinstance(block, Table):
            for row in block.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))

    raw = "\n\n".join(parts)
    raw_pages = _fixed_size_pseudo_pages(raw)
    pages = []
    for i, chunk in enumerate(raw_pages, start=1):
        text = chunk.strip()
        if text:
            pages.append(
                Page(
                    doc_id=doc["doc_id"],
                    title=doc["title"],
                    doc_type=doc["doc_type"],
                    effective_date=doc["effective_date"],
                    supersedes=doc["supersedes"],
                    page_number=i,
                    text=text,
                    faculty=doc.get("faculty"),
                    approx_pages=True,
                )
            )
    return pages


def parse_pdf(doc: dict) -> list[Page]:
    path = Path(doc["path"])
    fitz_doc = pymupdf.open(str(path))
    pages = []
    for i, page in enumerate(fitz_doc, start=1):
        text = page.get_text().strip()
        ocr = False
        if not text and doc.get("requires_ocr"):
            text = _ocr_page(path, i).strip()
            ocr = True
        if text:
            pages.append(
                Page(
                    doc_id=doc["doc_id"],
                    title=doc["title"],
                    doc_type=doc["doc_type"],
                    effective_date=doc["effective_date"],
                    supersedes=doc["supersedes"],
                    page_number=i,
                    text=text,
                    ocr=ocr,
                    faculty=doc.get("faculty"),
                )
            )
    return pages


def parse_document(doc: dict) -> list[Page]:
    path = Path(doc["path"])
    suffix = path.suffix.lower()
    if suffix == ".txt":
        return parse_txt(doc)
    if suffix == ".docx":
        return parse_docx(doc)
    return parse_pdf(doc)


def ingest_all(verbose: bool = True) -> list[Page]:
    all_pages: list[Page] = []
    for doc in DOCUMENTS:
        path = Path(doc["path"])
        if not path.exists():
            if verbose:
                print(f"[skip] {doc['doc_id']}: file not found at {path}")
            continue
        pages = parse_document(doc)
        all_pages.extend(pages)
        if verbose:
            n_ocr = sum(1 for p in pages if p.ocr)
            print(
                f"[ok]   {doc['doc_id']:30s} pages={len(pages):3d}  "
                f"ocr_pages={n_ocr:3d}  chars={sum(len(p.text) for p in pages):7d}"
            )
    return all_pages


def main():
    pages = ingest_all()
    out_path = DATA_PROCESSED / "pages.jsonl"
    with open(out_path, "w", encoding="utf-8") as f:
        for p in pages:
            f.write(json.dumps(asdict(p), ensure_ascii=False) + "\n")
    print(f"\nWrote {len(pages)} pages -> {out_path}")


if __name__ == "__main__":
    main()
